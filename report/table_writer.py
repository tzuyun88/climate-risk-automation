import pandas as pd
from docx import Document


def insert_df_as_table(doc: Document,
                       placeholder: str,
                       df: pd.DataFrame,
                       position: str = "before",
                       clear_placeholder: bool = True,
                       allow_empty: bool = False) -> bool:
    """
    在 Word 文件的 {{placeholder}} 位置插入 DataFrame 表格。

    Parameters
    ----------
    position : "before" | "after"
        表格插入在 placeholder 段落的前方或後方。
    allow_empty : bool
        為 True 時允許插入空白表格（僅含表頭）。
    """
    if df is None or (df.empty and not allow_empty):
        print(f"⚠ DataFrame 為空，無法在 '{placeholder}' 插入表格")
        return False

    # 建立表格（第 0 列為表頭）
    rows = df.shape[0] + 1
    cols = df.shape[1]
    table = doc.add_table(rows=rows, cols=cols)

    try:
        table.style = "Table Grid"
    except Exception:
        pass

    # 填入表頭
    for j, col in enumerate(df.columns):
        table.cell(0, j).text = str(col)

    # 填入資料列
    for i in range(df.shape[0]):
        for j in range(df.shape[1]):
            val = df.iloc[i, j]
            table.cell(i + 1, j).text = "" if pd.isna(val) else str(val)

    # 找到 placeholder 並插入表格
    placed = False
    for para in doc.paragraphs:
        if placeholder in para.text:
            if position == "before":
                para._p.addprevious(table._tbl)
            else:
                para._p.addnext(table._tbl)

            if clear_placeholder:
                para.text = para.text.replace(placeholder, "")

            placed = True
            break

    if not placed:
        print(f"[WARNING] 找不到 placeholder：'{placeholder}'，表格已新增但位置可能不對")

    return placed