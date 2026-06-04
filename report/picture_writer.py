import os
from docx import Document
from docx.shared import Inches
from typing import Optional


# ---------- 內部輔助函式 ----------

def _collect_all_paragraphs(doc: Document):
    """收集文件中所有段落（本文 + 表格儲存格 + 頁首頁尾）。"""
    paras = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paras.extend(cell.paragraphs)
    for section in doc.sections:
        paras.extend(section.header.paragraphs)
        paras.extend(section.footer.paragraphs)
    return paras


def _get_full_text(para) -> str:
    """取得段落中所有 run 組合後的完整文字。"""
    return "".join(run.text for run in para.runs)


# ---------- 對外介面 ----------

def find_picture_file(picture_dir: str, index: int) -> Optional[str]:
    """
    依 index 尋找對應圖片檔（picture1 ~ picture10）。
    支援常見大小寫與拼寫變體，以及常見圖片格式。
    """
    exts = [".png", ".jpg", ".jpeg", ".bmp", ".gif",
            ".tif", ".tiff", ".webp"]
    base_names = [
        f"picture{index}",  f"picture {index}",
        f"Picture{index}",  f"Picture {index}",
        f"pucture{index}",  f"pucture {index}",
        f"Pucture{index}",  f"Pucture {index}",
    ]
    for base in base_names:
        for ext in exts:
            candidate = os.path.join(picture_dir, base + ext)
            if os.path.exists(candidate):
                return candidate
    return None


def insert_picture_at_placeholder(doc: Document,
                                   placeholder: str,
                                   image_path: str,
                                   width_inches: float = 5.5) -> bool:
    """
    在 placeholder 段落下方插入圖片，並清除 placeholder 文字。
    """
    for para in _collect_all_paragraphs(doc):
        full_text = _get_full_text(para)
        if placeholder not in full_text and placeholder not in para.text:
            continue

        # 插入圖片段落
        img_para = doc.add_paragraph()
        run = img_para.add_run()
        run.add_picture(image_path, width=Inches(width_inches))
        para._p.addnext(img_para._p)

        # 清除 placeholder 文字
        if para.runs:
            combined = _get_full_text(para)
            cleaned = combined.replace(placeholder, "")
            for i, r in enumerate(para.runs):
                r.text = cleaned if i == 0 else ""
        else:
            para.text = para.text.replace(placeholder, "")

        print(f"✓ 插入圖片：{os.path.basename(image_path)} → {placeholder}")
        return True

    print(f"[WARNING] 找不到圖片 placeholder：'{placeholder}'")
    return False